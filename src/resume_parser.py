"""Resume text extraction from PDF and DOCX files."""

from pathlib import Path
from typing import Union

PathLike = Union[str, Path]


class ResumeParserError(Exception):
    """Base exception for resume parsing errors."""


class EmptyDocumentError(ResumeParserError):
    """Raised when a document contains no extractable text."""


class UnsupportedFormatError(ResumeParserError):
    """Raised when file extension is not supported."""


class CorruptedFileError(ResumeParserError):
    """Raised when file cannot be read or parsed."""


def extract_text_from_pdf(path: PathLike) -> str:
    """Extract text from a PDF file using PyMuPDF."""
    path = Path(path)
    if not path.exists():
        raise CorruptedFileError(f"File not found: {path}")

    try:
        import fitz  # PyMuPDF
    except ImportError as exc:
        raise CorruptedFileError("PyMuPDF is not installed.") from exc

    try:
        doc = fitz.open(str(path))
    except Exception as exc:
        raise CorruptedFileError(f"Cannot open PDF file: {exc}") from exc

    try:
        text_parts = []
        for page in doc:
            page_text = page.get_text("text")
            if page_text:
                text_parts.append(page_text)
        text = "\n".join(text_parts).strip()
    except Exception as exc:
        raise CorruptedFileError(f"Error reading PDF content: {exc}") from exc
    finally:
        doc.close()

    if not text:
        raise EmptyDocumentError("PDF contains no extractable text.")
    return text


def extract_text_from_docx(path: PathLike) -> str:
    """Extract text from a DOCX file using python-docx."""
    path = Path(path)
    if not path.exists():
        raise CorruptedFileError(f"File not found: {path}")

    try:
        from docx import Document
    except ImportError as exc:
        raise CorruptedFileError("python-docx is not installed.") from exc

    try:
        document = Document(str(path))
    except Exception as exc:
        raise CorruptedFileError(f"Cannot open DOCX file: {exc}") from exc

    try:
        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
        # Also extract from tables
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text.strip())
        text = "\n".join(paragraphs).strip()
    except Exception as exc:
        raise CorruptedFileError(f"Error reading DOCX content: {exc}") from exc

    if not text:
        raise EmptyDocumentError("DOCX contains no extractable text.")
    return text


def extract_resume_text(path: PathLike) -> str:
    """
    Extract resume text based on file extension.

    Supports: .pdf, .docx, .doc (attempts docx parser for .doc)
    """
    path = Path(path)
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return extract_text_from_pdf(path)
    if suffix in (".docx", ".doc"):
        return extract_text_from_docx(path)

    raise UnsupportedFormatError(
        f"Unsupported file format '{suffix}'. Please upload PDF or DOCX."
    )


def extract_text_from_bytes(content: bytes, filename: str) -> str:
    """Extract resume text from uploaded bytes (for Streamlit file upload)."""
    import tempfile

    suffix = Path(filename).suffix.lower()
    if suffix not in (".pdf", ".docx", ".doc"):
        raise UnsupportedFormatError(
            f"Unsupported file format '{suffix}'. Please upload PDF or DOCX."
        )

    with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
        tmp.write(content)
        tmp_path = tmp.name

    try:
        return extract_resume_text(tmp_path)
    finally:
        Path(tmp_path).unlink(missing_ok=True)
